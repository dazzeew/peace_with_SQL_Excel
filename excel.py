import openpyxl as o

from bd_operation import io_listing, month_list_pgl, month_list, storage

from datetime import datetime

import math

# month_list_reverse = {1:'Январь',2:'Февраль',3:'Март',4:'Апрель',5:'Май',6:'Июнь',7:'Июль',8:'Август',9:'Сентябрь',10:'Октябрь',11:'Ноябрь',12:'Декабрь'}

def calculate_days(day):
	return math.ceil(day.days / 31) + 1

def create_null_mass(date_f,day):
	itog_mass = []
	year = int(date_f[0] + date_f[1] + date_f[2] + date_f[3])
	month = int(date_f[4] + date_f[5])
	for i in range(day):
		itog_mass.append((0,month,year))
		if month == 12:
			year += 1
		month = (month + 1) % 13
		if month == 0:
			month = 1
	return itog_mass

def calculate_money(list):
	mass = io_listing(list)
	day = calculate_days(datetime.strptime(list[1],'%Y%m%d') - datetime.strptime(list[0],'%Y%m%d'))
	summ = 0
	c = 0
	year = int(list[0][0] + list[0][1] + list[0][2] + list[0][3])
	itog_mass = create_null_mass(list[0],day)
	result = []
	j = int(datetime.date(datetime.strptime(mass[0][2],'%Y-%m-%d')).strftime('%m'))
	for i in range(day):
		while ((int(mass[0][2][5] + mass[0][2][6]) == j) and (int(mass[0][2][0] + mass[0][2][1] + mass[0][2][2] + mass[0][2][3]) == year)):
			summ += mass[0][5]
			mass.pop(0)
			if len(mass) == 0:
				break
		if ((j == itog_mass[i][1]) and (year == itog_mass[i][2])):
			result.append((summ,'=DATE(' + str(year) + ',' + str(j) + ',1)'))
		else:
			result.append((0,'=DATE(' + str(itog_mass[i][2]) + ',' + str(itog_mass[i][1]) + ',1)'))
		if len(mass) == 0:
			break
		if j == 12:
			year += 1
		j = (j + 1) % 13
		if j == 0:
			j = 1 
		summ = 0
	return result


def calculate_date(list):
	date_list = []
	if (list[0] is False):
		if (list[1] is False):
			if(list[2] is False):
				date_list = [str(list[4]) + month_list_pgl[list[3]] + '01',str(list[6]) + month_list_pgl[list[5]] + '01']
			else:
				date_list = [str(list[4]) + month_list_pgl[list[3]] + '01',str((datetime.date(datetime.now())).strftime('%Y%m')) + '01']
		else:
			date_list = [str(list[4]) + month_list_pgl[list[3]] + '01',str(list[4]) + month_list_pgl[list[3]] + '01']
	else:
		date_list = [str((datetime.date(datetime.now())).strftime('%Y%m')) + '01',str((datetime.date(datetime.now())).strftime('%Y%m')) + '01']
	if (datetime.strptime(date_list[0],'%Y%m%d') > datetime.strptime(date_list[1],'%Y%m%d')):
		date_list = [date_list[1],date_list[0]]
	if (date_list[0] == date_list[1]):
		date_list.append(False)
	else:
		date_list.append(True)
	return date_list


def excel_generation(arg):
	if arg[3] is True:
		dir = arg[13] + '/' + 'Dios.xlsx'
	else:
		dir = arg[13] + '/' + arg[11] + '.xlsx'
	print(arg[13])
	wb = o.Workbook()
	if (arg[6] == 4):
		sheet = wb['Sheet']
		sheet.title = 'Склад'
		sheet.column_dimensions['A'].width = 20
		sheet.column_dimensions['B'].width = 20
		sheet.column_dimensions['C'].width = 20
		sheet.column_dimensions['D'].width = 20
		sheet['A1'] = 'Название'
		sheet['B1'] = 'Тип'
		sheet['C1'] = 'Покупная цена'
		sheet['D1'] = 'Количество'
		for row in storage():
			sheet.append(row)
	else:
		date_list = [arg[0],arg[1],arg[2],arg[7],arg[8],arg[9],arg[10]]
		date_list = calculate_date(date_list)
		list = [date_list[0],date_list[1],arg[17],arg[4],arg[5],arg[6],arg[12],arg[14],arg[15],arg[16]]
		itog_list = io_listing(list)
		list_cp = [('Компания','Продукт','Дата','Количество','Цена за 1 единицу','Итоговая стоимость')]
		if (arg[6] == 0):
			sheet = wb['Sheet']
			sheet.title = 'Список закупок'
			sheet.column_dimensions['A'].width = 20
			sheet.column_dimensions['B'].width = 20
			sheet.column_dimensions['C'].width = 10
			for row in list_cp:
				sheet.append(row)
			for row in itog_list:
				sheet.append(row)
		elif (arg[6] == 1):
			sheet = wb['Sheet']
			sheet.title = 'Список продаж'
			sheet.column_dimensions['A'].width = 20
			sheet.column_dimensions['B'].width = 20
			sheet.column_dimensions['C'].width = 10
			for row in list_cp:
				sheet.append(row)
			for row in itog_list:
				sheet.append(row)
		elif (arg[6] == 2):
			sheet = wb['Sheet']
			sheet.title = 'Чистая прибыль'
			sheet['A1'] = 'Прибыль с '
			sheet.column_dimensions['B'].width = 10
			sheet.column_dimensions['C'].width = 5
			sheet.column_dimensions['D'].width = 10
			sheet['B1'] = (datetime.strptime(date_list[0],'%Y%m%d')).strftime('%Y-%m-%d')
			sheet['C1'] = 'по'
			sheet['D1'] = (datetime.strptime(date_list[1],'%Y%m%d')).strftime('%Y-%m-%d')
			sheet['A2'] = str(itog_list[0][0]) + ' руб.'
			sheet['A3'] = 'Прибыль'
			sheet['B3'] = 'Месяц/Год'
			sheet['E1'] = 'Компания'
			sheet['F1'] = 'Продукт'
			if arg[4] is True:
				sheet['E2'] = 'Все'
			else:
				sheet['E2'] = arg[12]
			if arg[17] is True:
				sheet['F2'] = 'Все'
			else:
				sheet['F2'] = arg[15]
			if (date_list[0] != date_list[1]):
				value = 3 + calculate_days(datetime.strptime(list[1],'%Y%m%d') - datetime.strptime(list[0],'%Y%m%d'))
				list[5] = 1
				for row in calculate_money(list):
					sheet.append(row)
				chart = o.chart.LineChart()
				bar = o.chart.BarChart()
				bar.y_axis.title = 'Прибыль руб.'
				bar.x_axis.title = 'Месяц/Год'
				chart.y_axis.title = 'Прибыль руб.'
				chart.x_axis.title = 'Месяц/Год'
				data = o.chart.Reference(sheet, min_col=1, min_row=4, max_row=value)
				dates = o.chart.Reference(sheet,min_col=2,min_row=4,max_row=value)
				for i in range(value - 2):
					sheet["B" + str(4 + i)].number_format = 'mmm-yy'
				bar.add_data(data)
				bar.set_categories(dates)
				chart.add_data(data)
				chart.set_categories(dates)
				sheet.add_chart(chart,'D7')
				sheet.add_chart(bar,'N7')
		elif (arg[6] == 3):
			sheet = wb['Sheet']
			sheet.title = 'Чистые траты'
			sheet['A1'] = 'Траты с '
			sheet.column_dimensions['B'].width = 10
			sheet.column_dimensions['C'].width = 3
			sheet.column_dimensions['D'].width = 10
			sheet['B1'] = (datetime.strptime(date_list[0],'%Y%m%d')).strftime('%Y-%m-%d')
			sheet['C1'] = 'по'
			sheet['D1'] = (datetime.strptime(date_list[1],'%Y%m%d')).strftime('%Y-%m-%d')
			sheet['A2'] = str(itog_list[0][0]) + ' руб.'
			sheet['A3'] = 'Прибыль'
			sheet['B3'] = 'Месяц/Год'
			sheet['E1'] = 'Компания'
			sheet['F1'] = 'Продукт'
			if arg[4] is True:
				sheet['E2'] = 'Все'
			else:
				sheet['E2'] = arg[12]
			if arg[17] is True:
				sheet['F2'] = 'Все'
			else:
				sheet['F2'] = arg[15]
			if (date_list[0] != date_list[1]):
				value = 3 + calculate_days(datetime.strptime(list[1],'%Y%m%d') - datetime.strptime(list[0],'%Y%m%d'))
				list[5] = 0
				for row in calculate_money(list):
					sheet.append(row)
				chart = o.chart.LineChart()
				bar = o.chart.BarChart()
				bar.y_axis.title = 'Траты руб.'
				bar.x_axis.title = 'Месяц/Год'
				chart.y_axis.title = 'Траты руб.'
				chart.x_axis.title = 'Месяц/Год'
				data = o.chart.Reference(sheet, min_col=1, min_row=4, max_row=value)
				dates = o.chart.Reference(sheet,min_col=2,min_row=4,max_row=value)
				for i in range(value - 2):
					sheet["B" + str(4 + i)].number_format = 'mmm-yy'
				bar.add_data(data)
				bar.set_categories(dates)
				chart.add_data(data)
				chart.set_categories(dates)
				chart.width = 17
				bar.width = 17
				sheet.add_chart(chart,'D7')
				sheet.add_chart(bar,'N7')
	wb.save(dir)